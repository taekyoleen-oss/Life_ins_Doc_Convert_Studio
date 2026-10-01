"""
엑셀 검산 파일 — 기본 상품 종신보험(암진단 포함)의 산출방법서 식을 엑셀 수식으로 적는다.

  입력      기초율·사업비·계약 한 점(남 40세 · 20년납 · 월납) — 값을 바꾸면 수식이 다시 계산된다
  위험률    연령 × 사망률 · 80% 이상 장해율 · 암발생률 (계약 성별)
  기수_*    담보마다 t · 나이 · q · 탈퇴 · f(납입면제) · Q · l · l′ · D · D′ · C · S · N · N′ · V (모두 수식)
  보험료    담보별 PVB · N* · P · P_base · G · 10만원당 G · 월보험료 — 앱(자유설계보험) 계산값과 차이

읽는 파일: samples/09_종신보험(암진단포함)_MethodSpec.json (Studio 기본 상품 · tests/default-product.test.ts)
          samples/09_종신보험(암진단포함)_계산결과.json (자유설계보험 계산 · flexible tests/ui/default-product.test.ts)
  python scripts/make-verify-xlsx.py      → samples/09_종신보험(암진단포함)_검산.xlsx
  powershell -File scripts/check-verify-xlsx.ps1   → 엑셀로 다시 계산해 차이를 확인하고 계산된 값까지 저장
"""
import json

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

BASE = "samples/09_종신보험(암진단포함)"
spec = json.load(open(f"{BASE}_MethodSpec.json", encoding="utf-8"))
res = json.load(open(f"{BASE}_계산결과.json", encoding="utf-8"))
sex, x, m, k = res["contract"]["sex"], res["contract"]["age"], res["contract"]["payYears"], res["contract"]["freq"]
rates = {r["id"]: r for r in spec["rates"]}
exp = {e["symbol"]: e.get("rate", e.get("times")) for e in spec["expenses"]}
basis = spec["basis"]

HEAD = PatternFill("solid", fgColor="DCE6F1")
INPUT = PatternFill("solid", fgColor="FFF4D6")
APP = PatternFill("solid", fgColor="E2EFDA")
BOLD = Font(bold=True)
THIN = Border(bottom=Side(style="thin", color="BBBBBB"))
wb = openpyxl.Workbook()


def name(n, ref):
    dn = DefinedName(n, attr_text=ref)
    try:
        wb.defined_names[n] = dn            # openpyxl 3.1
    except TypeError:
        wb.defined_names.append(dn)         # 옛 판


# ── 설명 ────────────────────────────────────────────────────────────────────
ws = wb.active
ws.title = "설명"
lines = [
    f"{spec['meta']['productName']} — 산출방법서 식 엑셀 검산",
    "",
    f"계약 한 점: {'남' if sex == 'M' else '여'} {x}세 · 보험료 납입기간 {m}년 · 월납(k={k}) · 보험기간은 담보 만기(종신 110세)",
    "담보: 사망·80% 이상 장해(사망형, 1억원, 110세) · 암 진단(진단형, 사망보험금의 50% = 5천만원, 100세, 90일 면책 → 첫해 급부 × (1 − 3/12))",
    "납입면제: 80% 이상 장해 · 암 진단 시 이후 보험료 면제. 사망은 계약 소멸(탈퇴). 납입면제 사유가 그 담보의 탈퇴 사유이기도 하면 다시 빼지 않는다",
    "   사망·80% 이상 장해: 탈퇴 = 사망 + 장해, 납입면제 f = 암 발생   /   암 진단: 탈퇴 = 사망 + 암 발생, 납입면제 f = 80% 이상 장해",
    "",
    "시트",
    "  입력      기초율·사업비·계약 (노란 칸을 바꾸면 모든 시트가 다시 계산된다)",
    "  위험률    연령 × 위험률 (계약 성별 — 공개 기본 위험률: 사망률 · 80% 이상 장해율 · 암발생률)",
    "  기수_*    담보마다 계산기수와 연말 책임준비금 V — 모두 수식. 초록 열 = 앱(자유설계보험) 값, 그 옆 = 차이",
    "  보험료    담보별 보험료와 앱 값의 차이 — 차이가 0 이면 산출방법서 식 = 앱 계산",
    "",
    "식 (산출방법서 3~5절과 같다)",
    "  탈퇴율 Q = q + r − q·r/2 ,  l_{t+1} = l_t·(1 − Q) ,  질병 F = 1 − (1 − r)(1 − f) ,  Q′ = q + F − q·F/2 ,  l′_{t+1} = l′_t·(1 − Q′)",
    "  (질병끼리는 곱으로, 사망과는 겹치는 부분을 절반으로 결합한다 · 보장기간 n = 만기 나이 − 가입나이, 종신(110세)은 + 1)",
    "  D = l·v^t ,  D′ = l′·v^t ,  C = l·g·v^{t+½} (사망형 g = Q, 진단형 g = 그 발생률) ,  N = Σ_{u≥t} D ,  N′ = Σ_{u≥t} D′",
    "  PVB = Σ S·C ,  N* = k·[(N′_x − N′_{x+m}) − (k−1)/(2k)·(D′_x − D′_{x+m})] ,  P = PVB/N* ,  P_base = PVB/(N′_x − N′_{x+min(n,20)})",
    "  G = [P + (α_S + α_P·P_base)·D′_x/N* + β_S/k + β′·(N_{x+m} − N_{x+n})/N*] / (1 − β_G − γ) ,  10만원당 G = ROUND(ROUND(G, 6)×100,000, 0)",
    "  V_t = [Σ_{u≥t} S·C + β′·(N_{x+max(t,m)} − N_{x+n}) − P_β·(N′_{x+t} − N′_{x+m})·[t≤m]] / D_{x+t} ,  P_β = [PVB + β′·(N_{x+m} − N_{x+n})] / (N′_x − N′_{x+m})",
    "",
    "행 수는 이 계약(40세)에 맞춰져 있다 — 가입나이를 바꾸려면 앱에서 다시 만든다(python scripts/make-verify-xlsx.py).",
]
for i, t in enumerate(lines, 1):
    ws.cell(i, 1, t)
ws["A1"].font = Font(bold=True, size=14)
ws.column_dimensions["A"].width = 150

# ── 입력 ────────────────────────────────────────────────────────────────────
wi = wb.create_sheet("입력")
wi.append(["항목", "값", "이름(수식에서)", "비고"])
inputs = [
    ("적용이율 i", basis["interest"], "이율", "예정이율"),
    ("현가율 v", "=1/(1+이율)", "현가율", "= 1/(1+i)"),
    ("가입나이 x", x, "가입나이", "남" if sex == "M" else "여"),
    ("보험료 납입기간 m", m, "납입기간", "년"),
    ("납입주기별 계수 k", k, "납입주기", "월납 12"),
    ("보험기간(계약)", "=MAX(IF(기수_사망장해!B3>=110,기수_사망장해!B3+1,기수_사망장해!B3),IF(기수_암진단!B3>=110,기수_암진단!B3+1,기수_암진단!B3))-가입나이", "보험기간", "담보 만기 중 가장 늦은 것"),
    ("α_S 신계약비(가입금액)", exp["α_S"], "알파S", "보험가입금액 비례"),
    ("α_P 신계약비(기준연납순보험료 배수)", exp["α_P"], "알파P", "보장기간 20년 미만이면 × n/20"),
    ("β_S 유지비(납입중)", exp["β_S"], "베타S", "매년 보험가입금액 — 1회 납입당 β_S/k"),
    ("β_G 유지비(영업보험료)", exp["β_G"], "베타G", ""),
    ("β′ 납입후 유지비", exp["β′"], "베타후", "매년 보험가입금액"),
    ("γ 수금비", exp["γ"], "감마", "영업보험료"),
]
for i, (lab, v, nm, note) in enumerate(inputs, 2):
    wi.cell(i, 1, lab); wi.cell(i, 2, v); wi.cell(i, 3, nm); wi.cell(i, 4, note)
    if not (isinstance(v, str) and v.startswith("=")):
        wi.cell(i, 2).fill = INPUT
    name(nm, f"입력!$B${i}")
for c in "ABCD":
    wi[f"{c}1"].fill = HEAD; wi[f"{c}1"].font = BOLD
wi.column_dimensions["A"].width = 34; wi.column_dimensions["B"].width = 14; wi.column_dimensions["C"].width = 14; wi.column_dimensions["D"].width = 40

# ── 위험률 ──────────────────────────────────────────────────────────────────
wr = wb.create_sheet("위험률")
cols = [("q", "사망률 q"), ("r80", "80% 이상 장해율"), ("rc", "암발생률")]
wr.append(["연령"] + [f"{rates[i]['name']} ({'남' if sex == 'M' else '여'})" for i, _ in cols])
tabs = [dict(zip(rates[i]["tables"][sex]["ages"], rates[i]["tables"][sex]["values"])) for i, _ in cols]
ages = sorted(set().union(*[t.keys() for t in tabs]))
for a in ages:
    wr.append([a] + [t.get(a) for t in tabs])
for c in "ABCD":
    wr[f"{c}1"].fill = HEAD; wr[f"{c}1"].font = BOLD
    wr.column_dimensions[c].width = 22 if c != "A" else 8
for row in wr.iter_rows(min_row=2, min_col=2, max_col=4):
    for cell in row:
        cell.number_format = "0.000000"; cell.fill = INPUT
wr.freeze_panes = "B2"
RATE_COL = {"q": "B", "r80": "C", "rc": "D"}      # 위험률 시트 열 — 행 = 나이 + 2


# ── 담보별 계산기수 ─────────────────────────────────────────────────────────
def coverage(sheet, ben, other, waiver, event_is_q, app):
    """other = 두 번째 탈퇴 위험률 id, waiver = 납입면제 f 위험률 id, event_is_q = 사망형(급부 = 탈퇴 전부)"""
    w = wb.create_sheet(sheet)
    n = app["n"]
    R0 = 16                                   # 보험료 칸(2~13행) 아래 — 15행이 표 머리
    last = R0 + n
    rng = lambda c: f"${c}${R0}:${c}${last}"
    months = round(ben.get("waitDays", 0) / 30.4) if ben.get("waitDays") else 0
    rows = [
        ("담보", ben["name"]), ("보장금액", app["amount"]), ("만기 나이", ben["endAge"]),
        ("보장기간 n", "=MIN(보험기간,IF(B3>=110,B3+1,B3)-가입나이)"), ("납입기간 m", "=MIN(납입기간,B4)"),
        ("면책(개월)", months), ("첫해 급부 배율", "=MAX(0,1-B6/12)"),
        ("탈퇴 사유", f"사망 + {rates[other]['name']}"), ("납입면제 f", rates[waiver]["name"]),
        ("급부 g", "탈퇴 전부 (사망형 — g = Q)" if event_is_q else f"{rates[other]['name']} (진단형)"),
    ]
    for i, (a, b) in enumerate(rows, 1):
        w.cell(i, 1, a).font = BOLD; w.cell(i, 2, b)
        if a in ("보장금액", "만기 나이", "면책(개월)"):
            w.cell(i, 2).fill = INPUT
    # 보험료 — 식 · 엑셀 값 · 앱 값 · 차이
    w["D1"], w["E1"], w["F1"], w["G1"] = "항목", "엑셀(수식)", "앱(자유설계보험)", "차이"
    for c in "DEFG":
        w[f"{c}1"].fill = HEAD; w[f"{c}1"].font = BOLD
    prem = [
        ("PVB = Σ S·C", f"=SUMPRODUCT($L${R0}:$L${last - 1},$K${R0}:$K${last - 1})", app["pvb"]),
        ("N* (연납 환산 납입기수)", f"=납입주기*(($N${R0}-INDEX({rng('N')},B5+1))-(납입주기-1)/(2*납입주기)*($J${R0}-INDEX({rng('J')},B5+1)))", app["nStar"]),
        ("P = PVB/N*", "=E2/E3", app["net"]),
        ("P_base (기준연납순보험료)", f"=E2/($N${R0}-INDEX({rng('N')},MIN(B4,20)+1))", app["base"]),
        ("α_P 적용 = α_P·min(n,20)/20", "=알파P*MIN(B4,20)/20", None),
        ("신계약비 부가 (α_S+α_P·P_base)·D′x/N*", f"=(알파S+E6*E5)*$J${R0}/E3", None),
        ("β_S 부가 = β_S/k", "=베타S/납입주기", None),
        ("β′ 부가 = β′·(N_m − N_n)/N*", f"=베타후*(INDEX({rng('M')},B5+1)-$M${last})/E3", None),
        ("G (영업보험료, 1원당)", "=(E4+E7+E8+E9)/(1-베타G-감마)", app["gross"]),
        ("10만원당 G (1원당 6자리 → 10만원당 반올림)", "=ROUND(ROUND(E10,6)*100000,0)", app["gross100k"]),
        ("월보험료 = 10만원당 G × 보장금액/10만", "=E11*B2/100000", app["monthlyGross"]),
        ("P_β (준비금용 순보험료)", f"=(E2+베타후*(INDEX({rng('M')},B5+1)-$M${last}))/($N${R0}-INDEX({rng('N')},B5+1))", app["pBeta"]),
    ]
    for i, (lab, f, v) in enumerate(prem, 2):
        w.cell(i, 4, lab); w.cell(i, 5, f)
        if v is not None:
            w.cell(i, 6, v).fill = APP
            w.cell(i, 7, f"=E{i}-F{i}")
        w.cell(i, 5).number_format = w.cell(i, 6).number_format = "0.000000000" if i not in (11, 12) else "#,##0"
        w.cell(i, 7).number_format = "0.0E+00"
    # 표
    head = ["t", "나이", "q 사망률", f"탈퇴2 {rates[other]['name']}", f"f {rates[waiver]['name']}", "Q 탈퇴율", "l 유지자수", "l′ 납입자수",
            "D", "D′", "C", "S", "N", "N′", "V (1원당)", "V 10만원당", "앱 V 10만원당", "차이"]
    for j, h in enumerate(head, 1):
        c = w.cell(R0 - 1, j, h); c.fill = APP if h.startswith("앱") else HEAD; c.font = BOLD; c.alignment = Alignment(wrap_text=True)
    ev = "F" if event_is_q else "D"
    for t in range(n + 1):
        r = R0 + t
        age_row = f"B{r}+2"
        w.cell(r, 1, t)
        w.cell(r, 2, f"=가입나이+A{r}")
        w.cell(r, 3, f"=INDEX(위험률!$B:$B,{age_row})")
        w.cell(r, 4, f"=INDEX(위험률!${RATE_COL[other]}:${RATE_COL[other]},{age_row})")
        w.cell(r, 5, f"=INDEX(위험률!${RATE_COL[waiver]}:${RATE_COL[waiver]},{age_row})")
        w.cell(r, 6, f"=MIN(1,C{r}+D{r}-C{r}*D{r}/2)")
        w.cell(r, 7, 100000 if t == 0 else f"=G{r - 1}*MAX(0,1-F{r - 1})")
        w.cell(r, 8, 100000 if t == 0 else f"=H{r - 1}*MAX(0,1-MIN(1,C{r - 1}+(1-(1-D{r - 1})*(1-E{r - 1}))-C{r - 1}*(1-(1-D{r - 1})*(1-E{r - 1}))/2))")
        w.cell(r, 9, f"=G{r}*현가율^A{r}")
        w.cell(r, 10, f"=H{r}*현가율^A{r}")
        w.cell(r, 11, f"=G{r}*{ev}{r}*현가율^(A{r}+0.5)")
        w.cell(r, 12, f"=IF(A{r}=0,$B$7,1)")
        w.cell(r, 13, f"=SUM(I{r}:$I${last})")
        w.cell(r, 14, f"=SUM(J{r}:$J${last})")
        w.cell(r, 15, f"=IF(I{r}<=0,0,(IF(A{r}<$B$4,SUMPRODUCT(L{r}:$L${last - 1},K{r}:$K${last - 1}),0)"
                      f"+베타후*(INDEX({rng('M')},MAX(A{r},$B$5)+1)-$M${last})-IF(A{r}<=$B$5,$E$13*(N{r}-INDEX({rng('N')},$B$5+1)),0))/I{r})")
        w.cell(r, 16, f"=ROUND(O{r}*100000,0)")
        w.cell(r, 17, app["reserve100k"][t]).fill = APP
        w.cell(r, 18, f"=P{r}-Q{r}")
        for j, fmt in zip(range(3, 19), ["0.000000"] * 4 + ["#,##0.000"] * 2 + ["#,##0.000"] * 3 + ["0.00"] + ["#,##0.000"] * 2 + ["0.0000000"] + ["#,##0"] * 3):
            w.cell(r, j).number_format = fmt
    w.freeze_panes = f"C{R0}"
    widths = [5, 6, 11, 13, 13, 11, 12, 12, 12, 12, 12, 6, 13, 13, 12, 11, 11, 7]
    for j, wd in enumerate(widths, 1):
        w.column_dimensions[openpyxl.utils.get_column_letter(j)].width = wd
    w.column_dimensions["D"].width = 34; w.column_dimensions["E"].width = 16; w.column_dimensions["F"].width = 16
    return w


b1, b2 = spec["benefits"]
a1, a2 = res["coverages"]
coverage("기수_사망장해", b1, "r80", "rc", True, a1)
coverage("기수_암진단", b2, "rc", "r80", False, a2)

# ── 보험료 요약 ─────────────────────────────────────────────────────────────
wp = wb.create_sheet("보험료", 1)
head = ["담보", "n", "m", "PVB", "N*", "P", "P_base", "G (1원당)", "10만원당 G (엑셀)", "10만원당 G (앱)", "차이", "월보험료 (엑셀)", "월보험료 (앱)", "차이", "준비금 차이(최대)"]
wp.append(head)
for i, (sh, a) in enumerate((("기수_사망장해", a1), ("기수_암진단", a2)), 2):
    wp.append([f"='{sh}'!B1", f"='{sh}'!B4", f"='{sh}'!B5", f"='{sh}'!E2", f"='{sh}'!E3", f"='{sh}'!E4", f"='{sh}'!E5", f"='{sh}'!E10",
               f"='{sh}'!E11", a["gross100k"], f"=I{i}-J{i}", f"='{sh}'!E12", a["monthlyGross"], f"=L{i}-M{i}",
               f"=MAX(MAX('{sh}'!R16:R{16 + a['n']}),-MIN('{sh}'!R16:R{16 + a['n']}))"])
    wp.cell(i, 10).fill = APP; wp.cell(i, 13).fill = APP
wp.append(["합계", None, None, None, None, None, None, None, "=SUM(I2:I3)", "=SUM(J2:J3)", "=I4-J4", "=SUM(L2:L3)", res["monthlyGross"], "=L4-M4", "=MAX(O2:O3)"])
wp.cell(4, 13).fill = APP
for c in range(1, len(head) + 1):
    wp.cell(1, c).fill = HEAD; wp.cell(1, c).font = BOLD; wp.cell(1, c).alignment = Alignment(wrap_text=True)
    wp.cell(4, c).font = BOLD; wp.cell(4, c).border = THIN
    wp.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 14
wp.column_dimensions["A"].width = 22
for r in (2, 3):
    for c, fmt in zip(range(4, 9), ["#,##0.0000", "#,##0.00", "0.0000000000", "0.0000000000", "0.0000000000"]):
        wp.cell(r, c).number_format = fmt
for r in (2, 3, 4):
    for c in (9, 10, 11, 12, 13, 14, 15):
        wp.cell(r, c).number_format = "#,##0"
wp.append([])
wp.append(["차이 열이 모두 0 이면 산출방법서의 식을 엑셀 수식으로 계산한 보험료·준비금이 자유설계보험의 계산과 같다. (기수_* 시트의 R열 = 연도별 준비금 차이)"])

out = f"{BASE}_검산.xlsx"
wb.save(out)
print(out)
