"""
사내 위험률 모음(엑셀) → private/rate-library.json — 이 PC 에서만 쓰는 "기본 위험률 모음".

원본 첫 줄이 "사내 자료 — 외부 반출 금지" 이므로 결과는 private/ (깃·배포 제외)에만 둔다.
앱은 predev/prebuild 의 scripts/copy-public.mjs 가 public/rate-library.json 으로 복사한 것을 읽는다(그 파일도 깃 제외).
다른 PC·배포본에서는 이 파일이 없어 공개 기본 위험률만 보인다.

  python scripts/import-rate-library.py ["<위험률_모음_대표.xlsx 경로>"]
"""
import datetime
import json
import os
import sys

import openpyxl

DEFAULT = r"C:\Users\tklee\OneDrive - 코리안리재보험\위험률 산출과정 정리(클로드)\08_위험률모음\위험률_모음_대표.xlsx"
src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
wb = openpyxl.load_workbook(src, data_only=True, read_only=True)

# 00_목록 — ID 마다 분류·회사 근거·원본 위치
meta = {}
for r in wb["00_목록"].iter_rows(min_row=6, values_only=True):
    if not r[1]:
        continue
    meta[r[1]] = {
        "group": r[2], "category": r[3], "name": r[4], "detail": r[5], "sex": {"남": "M", "여": "F"}.get(r[6]),
        "company": r[7], "basis": r[8], "range": r[9],
        "origin": " · ".join(str(x) for x in (r[13], r[14], r[15]) if x),
    }

rates = []
for ws in wb.worksheets:
    if not (ws.title.startswith("결과") or ws.title.startswith("입력")):
        continue
    rows = list(ws.iter_rows(values_only=True))
    ids = rows[2][1:]                                   # 3행 = ID
    for j, rid in enumerate(ids, start=1):
        if not rid or rid not in meta:
            continue
        ages, values = [], []
        for row in rows[9:]:                            # 10행부터 연령(0~114) × 값
            a, v = row[0], row[j] if j < len(row) else None
            if isinstance(a, (int, float)) and isinstance(v, (int, float)):
                ages.append(int(a))
                values.append(float(v))
        if ages:
            rates.append({"id": rid, **{k: v for k, v in meta[rid].items() if v not in (None, "", "-")}, "ages": ages, "values": values})

out = {
    "title": "위험률 모음(대표)",
    "notice": "사내 자료 — 외부 반출 금지. 이 PC 의 private/ 에만 둔다(깃·배포 제외).",
    "source": os.path.basename(src),
    "createdAt": datetime.datetime.now().isoformat(timespec="seconds"),
    "rates": rates,
}
os.makedirs("private", exist_ok=True)
with open("private/rate-library.json", "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
print(f"private/rate-library.json: 위험률 {len(rates)}개 ({os.path.getsize('private/rate-library.json') // 1024}KB)")
